"""Document building blocks without a database: terbilang, formatting, PDF and XLSX rendering."""

from __future__ import annotations

import io
from datetime import UTC, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.modules.documents import formatting as fmt
from app.modules.documents.render_billing import render_invoice_pdf, render_quotation_pdf
from app.modules.documents.render_pdf import render_rab_pdf
from app.modules.documents.render_workorders import render_bast_pdf, render_spk_pdf
from app.modules.documents.render_xlsx import render_rab_xlsx
from app.modules.documents.snapshot import canonical_json, snapshot_sha256
from app.modules.documents.terbilang import rupiah_in_words, terbilang
from tests.document_samples import (
    bast_snapshot,
    invoice_snapshot,
    many_items,
    quotation_snapshot,
    sample_snapshot,
    spk_snapshot,
)

PREVIEW = {"preview": True, "number": None, "issued_at": None, "issued_by": None}


def pdf_text(content: bytes) -> list[str]:
    # Whitespace-normalized: extraction breaks lines wherever ReportLab positions a text run.
    return [" ".join(page.extract_text().split()) for page in PdfReader(io.BytesIO(content)).pages]


# --------------------------------------------------------------------------- terbilang


@pytest.mark.parametrize(
    ("n", "words"),
    [
        (0, "nol"),
        (1, "satu"),
        (10, "sepuluh"),
        (11, "sebelas"),
        (12, "dua belas"),
        (19, "sembilan belas"),
        (20, "dua puluh"),
        (21, "dua puluh satu"),
        (100, "seratus"),
        (101, "seratus satu"),
        (110, "seratus sepuluh"),
        (115, "seratus lima belas"),
        (999, "sembilan ratus sembilan puluh sembilan"),
        (1000, "seribu"),
        (1001, "seribu satu"),
        (1100, "seribu seratus"),
        (2000, "dua ribu"),
        (11_000, "sebelas ribu"),
        (100_000, "seratus ribu"),
        (1_000_000, "satu juta"),
        (1_200_000, "satu juta dua ratus ribu"),
        (
            10_240_776,
            "sepuluh juta dua ratus empat puluh ribu tujuh ratus tujuh puluh enam",
        ),
        (1_000_000_000, "satu miliar"),
        (2_500_000_000_000, "dua triliun lima ratus miliar"),
        (1_000_000_000_000_000, "satu kuadriliun"),
    ],
)
def test_terbilang(n: int, words: str) -> None:
    assert terbilang(n) == words


def test_terbilang_limits() -> None:
    assert terbilang(999_999_999_999_999_999).startswith(
        "sembilan ratus sembilan puluh sembilan kuadriliun"
    )
    with pytest.raises(ValueError):
        terbilang(-1)
    with pytest.raises(ValueError):
        terbilang(10**18)


def test_rupiah_in_words() -> None:
    assert rupiah_in_words("1000.00") == "Seribu rupiah"
    assert rupiah_in_words("0.50") == "Nol rupiah lima puluh sen"
    assert rupiah_in_words(Decimal("10240776.75")) == (
        "Sepuluh juta dua ratus empat puluh ribu tujuh ratus tujuh puluh enam rupiah "
        "tujuh puluh lima sen"
    )


# --------------------------------------------------------------------------- formatting


def test_money_and_numbers_never_go_through_float() -> None:
    assert fmt.money("1250000.00") == "Rp 1.250.000"
    assert fmt.money("1250000.50") == "Rp 1.250.000,50"
    assert fmt.money("0.00") == "Rp 0"
    assert fmt.money(None) == "-"
    assert fmt.money("123.40", "USD") == "USD 123,40"
    assert fmt.money("999999999999999999.99") == "Rp 999.999.999.999.999.999,99"
    assert fmt.unit_price("33.3350") == "Rp 33,335"
    assert fmt.unit_price("2767777.5000") == "Rp 2.767.777,50"
    assert fmt.unit_price("65000.0000") == "Rp 65.000"
    assert fmt.quantity("2.5000") == "2,5"
    assert fmt.quantity("1000.0000") == "1.000"
    assert fmt.percent("11.00") == "11"
    assert fmt.percent("2.50") == "2,5"
    assert fmt.number("-1500.5", min_places=2) == "-1.500,50"


def test_dates_are_shown_in_jakarta_time() -> None:
    assert fmt.tanggal("2026-10-19") == "19 Oktober 2026"
    late_utc = datetime(2026, 10, 18, 20, 0, tzinfo=UTC)  # 03.00 WIB the next day
    assert fmt.tanggal(late_utc) == "19 Oktober 2026"
    assert fmt.tanggal_waktu(late_utc.isoformat()) == "19 Oktober 2026 pukul 03.00 WIB"
    assert fmt.tanggal(None) == "-"


# --------------------------------------------------------------------------- snapshot hash


def test_snapshot_hash_is_canonical() -> None:
    a = sample_snapshot()
    b = sample_snapshot()
    b["company"] = dict(reversed(list(b["company"].items())))  # key order must not matter
    assert canonical_json(a) == canonical_json(b)
    assert snapshot_sha256(a) == snapshot_sha256(b)
    b["totals"]["grand_total"] = "10240776.76"
    assert snapshot_sha256(a) != snapshot_sha256(b)


# --------------------------------------------------------------------------- PDF


def test_issued_pdf_shows_the_snapshot_and_is_deterministic() -> None:
    snapshot = sample_snapshot()
    digest = "ab" * 32
    first = render_rab_pdf(snapshot, digest)
    assert first.startswith(b"%PDF")
    assert render_rab_pdf(snapshot, digest) == first  # same snapshot → same bytes

    text = "\n".join(pdf_text(first))
    for expected in (
        "RENCANA ANGGARAN BIAYA",
        "RAB-2026-0001/V1",
        "PT Contoh Teknik Nusantara",
        "NPWP 01.234.567.8-901.000",
        "CV Klien Makmur",
        "Kabel UTP Cat6",
        "Rp 1.311.500",
        "Subtotal Material",
        "Diskon 5%",
        "PPN 11%",
        "Rp 10.240.776,75",
        "Terbilang: Sepuluh juta",
        "Pembayaran 30% di muka.",
        "Dewi Approver",
        "Kode verifikasi: abababababababab",
        "Halaman 1 dari 1",
    ):
        assert expected in text, expected
    assert "PRATINJAU" not in text


def test_preview_pdf_is_watermarked_and_unnumbered() -> None:
    snapshot = sample_snapshot(
        document=PREVIEW, rab={"status_label": "Draft", "approved_by": None, "approved_at": None}
    )
    text = "\n".join(pdf_text(render_rab_pdf(snapshot)))
    assert "PRATINJAU: dokumen belum diterbitkan. Status versi: Draft" in text
    assert "PRATINJAU: bukan dokumen resmi" in text
    assert "RAB-2026-0001/V1" not in text
    assert "Diterbitkan oleh" not in text


def test_user_text_cannot_inject_paragraph_markup() -> None:
    # ReportLab paragraphs read <img src>, <a href> and <b>; user text must stay literal text.
    hostile = "<img src='C:/Windows/win.ini'/> & <b>tebal</b> <a href='http://x'>tautan</a>"
    items = sample_snapshot()["items"]
    items[0]["name"] = hostile
    items[0]["specification"] = "<para>spec</para>"
    snapshot = sample_snapshot(items=items, settings={"tagline": "<font size=40>BESAR</font>"})
    text = "\n".join(pdf_text(render_rab_pdf(snapshot, "0" * 64)))
    assert "<img src='C:/Windows/win.ini'/> & <b>tebal</b>" in text
    assert "<para>spec</para>" in text
    assert "<font size=40>BESAR</font>" in text


def test_long_rab_spans_pages_with_repeated_header() -> None:
    pages = pdf_text(render_rab_pdf(sample_snapshot(**many_items(120)), "0" * 64))
    assert len(pages) >= 3
    assert all("Uraian" in page for page in pages[:-1])  # header row repeats on every page
    assert f"Halaman 2 dari {len(pages)}" in pages[1]
    assert "(lanjutan)" in pages[1]
    assert "Item nomor 120" in "\n".join(pages)


def test_layout_options() -> None:
    hidden = "\n".join(
        pdf_text(render_rab_pdf(sample_snapshot(settings={"show_item_codes": False})))
    )
    assert "IT-01" not in hidden
    reader = PdfReader(
        io.BytesIO(render_rab_pdf(sample_snapshot(settings={"show_price_sources": True})))
    )
    box = reader.pages[0].mediabox
    assert box.width > box.height  # landscape to fit the source column
    assert "Toko Komputer Jaya" in " ".join(reader.pages[0].extract_text().split())


def test_characters_outside_the_font_do_not_break_rendering() -> None:
    items = sample_snapshot()["items"]
    items[0]["name"] = "Kabel 网络 ✓ → ≤ 5 m²"
    assert render_rab_pdf(sample_snapshot(items=items)).startswith(b"%PDF")


# --------------------------------------------------------------------------- XLSX


def test_xlsx_has_values_not_formulas_and_exact_numbers() -> None:
    items = sample_snapshot()["items"]
    items[0]["name"] = '=HYPERLINK("http://evil.test","klik")'
    items[1]["specification"] = "@SUM(A1:A2)"
    content = render_rab_xlsx(sample_snapshot(items=items), "ab" * 32)
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["RAB", "Info"]
    sheet = workbook["RAB"]
    cells = [cell for row in sheet.iter_rows() for cell in row if cell.value is not None]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    hostile = next(c for c in cells if c.value == '=HYPERLINK("http://evil.test","klik")')
    assert hostile.data_type == "s"
    assert hostile.quotePrefix is True
    assert any(c.value == "@SUM(A1:A2)" and c.data_type == "s" for c in cells)

    values = [c.value for c in cells]
    assert 1311500 in values  # item subtotal as an exact number
    assert Decimal(
        str(next(v for v in values if isinstance(v, float) and v > 10_000_000))
    ) == Decimal("10240776.75")
    assert sheet.freeze_panes is not None
    info = {row[0].value: row[1].value for row in workbook["Info"].iter_rows()}
    assert info["Nomor dokumen"] == "RAB-2026-0001/V1"
    assert info["Kode verifikasi (SHA-256 snapshot)"] == "ab" * 32


def test_xlsx_writes_unrepresentable_amounts_as_text() -> None:
    totals = dict(sample_snapshot()["totals"])
    totals["grand_total"] = "999999999999999999.99"  # beyond double precision
    workbook = load_workbook(io.BytesIO(render_rab_xlsx(sample_snapshot(totals=totals))))
    values = [c.value for row in workbook["RAB"].iter_rows() for c in row]
    assert "999999999999999999.99" in values


def test_xlsx_preview_is_marked() -> None:
    workbook = load_workbook(io.BytesIO(render_rab_xlsx(sample_snapshot(document=PREVIEW))))
    first_column = [row[0].value for row in workbook["RAB"].iter_rows(max_row=6)]
    assert any(isinstance(v, str) and v.startswith("PRATINJAU") for v in first_column)


# --------------------------------------------------------------------------- quotation & invoice


def test_quotation_pdf_shows_schedule_and_validity() -> None:
    snapshot = quotation_snapshot()
    content = render_quotation_pdf(snapshot, "ab" * 32)
    assert render_quotation_pdf(snapshot, "ab" * 32) == content
    text = " ".join(pdf_text(content))
    for expected in (
        "SURAT PENAWARAN HARGA",
        "PNW-2026-0001",
        "Berlaku sampai 18 November 2026",
        "Kepada CV Klien Makmur",
        "Jadwal pembayaran",
        "Uang muka 30% Rp 2.767.777,50 Rp 304.455,53 Rp 3.072.233,03",
        "Penawaran ini berlaku sampai 18 November 2026.",
    ):
        assert expected in text, expected
    assert "Toko Komputer Jaya" not in text  # price sources never appear on an offer


def test_invoice_pdf_payment_box_only_when_configured() -> None:
    text = " ".join(pdf_text(render_invoice_pdf(invoice_snapshot(), "cd" * 32)))
    for expected in (
        "INVOICE",
        "INV-2026-0001",
        "Jatuh tempo 2 November 2026",
        "NPWP 02.345.678.9-012.000",
        "Total tagihan Rp 3.072.233,03",
        "nomor rekening 123-456-7890",
    ):
        assert expected in text, expected
    bare = invoice_snapshot(
        payment_info={"bank_name": None, "account_number": None, "account_holder": None}
    )
    assert "Pembayaran" not in " ".join(pdf_text(render_invoice_pdf(bare)))


def test_billing_pdfs_escape_user_text() -> None:
    lines = invoice_snapshot()["lines"]
    lines[0]["description"] = "<img src='C:/x.png'/> & <b>x</b>"
    text = " ".join(pdf_text(render_invoice_pdf(invoice_snapshot(lines=lines))))
    assert "<img src='C:/x.png'/> & <b>x</b>" in text


# --------------------------------------------------------------------------- SPK & BAST


def test_spk_pdf_shows_parties_period_value_and_schedule() -> None:
    snapshot = spk_snapshot()
    content = render_spk_pdf(snapshot, "ef" * 32)
    assert render_spk_pdf(snapshot, "ef" * 32) == content
    text = " ".join(pdf_text(content))
    for expected in (
        "SURAT PERINTAH KERJA",
        "SPK-2026-0001",
        "Pemberi kerja CV Klien Makmur",
        "26 Oktober 2026 s.d. 25 November 2026",
        "Penawaran PNW-2026-0001",
        "Rincian volume pekerjaan",
        "Nilai kontrak Rp 10.240.776,75",
        "Jadwal pembayaran",
        "Budi Santoso Direktur",
        "Sari Wulandari Manajer Proyek",
        "Disetujui secara internal oleh Dewi Approver",
    ):
        assert expected in text, expected
    assert "Toko Komputer Jaya" not in text  # no prices or price sources on an SPK
    # Only what users wrote: no headings for empty terms/notes, no generated clauses.
    assert "Ketentuan" not in text
    assert "Catatan" not in text


def test_bast_pdf_shows_cumulative_volumes_without_assumed_conditions() -> None:
    text = " ".join(pdf_text(render_bast_pdf(bast_snapshot(), "01" * 32)))
    for expected in (
        "BERITA ACARA SERAH TERIMA",
        "BAST-2026-0001",
        "Tanggal serah terima 5 November 2026",
        "Serah terima sebagian",
        "SPK SPK-2026-0001",
        "Kabel UTP Cat6",
        "305 50 200 250",
        "Patch panel tambahan",
        "Yang menyerahkan",
        "Yang menerima",
    ):
        assert expected in text, expected
    assert "baik" not in text.lower()  # no "dalam keadaan baik" unless a user wrote it

    preview = bast_snapshot(document={**bast_snapshot()["document"], **PREVIEW})
    preview_text = " ".join(pdf_text(render_bast_pdf(preview)))
    assert "PRATINJAU" in preview_text
    assert "BAST-2026-0001" not in preview_text


def test_spk_and_bast_escape_user_text() -> None:
    hostile = "<img src='C:/x.png'/> & <b>x</b>"
    spk = spk_snapshot(spk={"scope": hostile, "client_signatory": {"name": hostile, "title": None}})
    assert hostile in " ".join(pdf_text(render_spk_pdf(spk)))
    lines = bast_snapshot()["lines"]
    lines[0]["condition"] = hostile
    assert hostile in " ".join(pdf_text(render_bast_pdf(bast_snapshot(lines=lines))))
