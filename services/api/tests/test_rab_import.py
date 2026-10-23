"""RAB item import: strict row parsing, preview → commit into a draft version, isolation."""

from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.modules.audit.models import AuditLog
from app.modules.identity.permissions import Role
from app.modules.rab.item_import import parse_file
from app.shared.spreadsheet import ImportFileError
from tests.factories import LoginAs, Tenant, make_project, make_user
from tests.test_rab_api import Rab, ok

HEADER = "kode;nama;kategori;jenis;volume;satuan;harga_satuan;sumber_harga\r\n"


def csv(*rows: str, header: str = HEADER) -> bytes:
    return (header + "".join(r + "\r\n" for r in rows)).encode("utf-8")


def xlsx(rows: list[list[Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- parsing


def test_rows_are_parsed_strictly_and_subtotals_come_from_the_calculator() -> None:
    parsed = parse_file(
        "item.csv",
        csv(
            "K-1;Semen 50 kg;Material;material;10;sak;65000;Toko A",
            ";Pekerja;Upah;tenaga kerja;2,5;OH;150000,50;",
            ";Bongkar;;jasa;1;ls;;",
            ";Besi;;;12;batang;4.500;",
            ";Nol;;;0;m;100;",
            ";Aneh;;robot;1;m;;",
            ";;;;1;m;1;",
            "K-1;Semen 50 kg;Material;material;10;sak;65000;Toko A",
        ),
    )
    rows = {r.row: r for r in parsed.rows}
    assert rows[2].valid and rows[2].data["subtotal"] == "650000.00"
    assert rows[2].data["quantity"] == "10" and rows[2].data["item_type"] == "material"
    assert rows[3].data["quantity"] == "2.5"
    assert rows[3].data["subtotal"] == "375001.25"  # 2.5 x 150000.50, ROUND_HALF_UP
    assert rows[3].data["item_type"] == "labor"
    assert rows[4].valid and rows[4].data["unit_price"] is None and rows[4].data["subtotal"] is None
    assert any("ambigu" in e for e in rows[5].errors)
    assert rows[6].errors == ["Volume: harus lebih dari 0."]
    assert any("tidak dikenal" in e for e in rows[7].errors)
    assert rows[8].errors == ["Nama: wajib diisi."]
    assert rows[9].errors == ["Duplikat dari baris 2."]


def test_aliases_formulas_and_headers() -> None:
    content = xlsx(
        [
            ["Uraian", "Kuantitas", "Satuan", "Harga", "Diskon"],
            ["Kabel", 305, "m", 4300.5, None],
            ["Rumus", "=2*3", "m", 1, None],
        ]
    )
    parsed = parse_file("item.xlsx", content)
    assert parsed.warnings == ["Kolom 'Diskon' tidak dikenal dan diabaikan."]
    assert parsed.rows[0].data["unit_price"] == "4300.5"
    assert any("rumus" in e for e in parsed.rows[1].errors)
    with pytest.raises(ImportFileError, match="volume"):
        parse_file("item.csv", csv(header="nama;satuan\r\n"))


# --------------------------------------------------------------------------- API


@pytest.fixture
def rab(db_session: Session, tenant: Tenant) -> Rab:
    project = make_project(db_session, tenant.org, currency="IDR")
    created = ok(
        tenant.browser.post(f"/api/v1/projects/{project.id}/rab", json={"title": "RAB Impor"}),
        201,
    )
    rab = Rab(tenant.browser, created)
    ok(rab.add(name="Item lama", quantity="1", unit="ls", unit_price="1000"), 201)
    return rab


def upload(
    rab: Rab, content: bytes, name: str = "item.csv", version: int = 1, browser: Any = None
) -> Any:
    return (browser or rab.browser).post(
        rab.url(version, "/imports"), files={"file": (name, content, "text/csv")}
    )


def test_template_downloads(tenant: Tenant) -> None:
    response = tenant.browser.get("/api/v1/rab/item-imports/template?format=xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Item RAB", "Petunjuk"]
    assert [c.value for c in workbook["Item RAB"][1]][:2] == ["kode", "nama"]
    text = tenant.browser.get("/api/v1/rab/item-imports/template?format=csv").content
    assert text.decode("utf-8-sig").startswith("kode;nama;kategori;jenis;spesifikasi;volume;satuan")
    # The template itself imports cleanly once filled in.
    workbook["Item RAB"].append([None, "Kabel", None, None, None, 10, "m", 4300])
    buffer = io.BytesIO()
    workbook.save(buffer)
    assert parse_file("t.xlsx", buffer.getvalue()).rows[0].valid


def test_preview_then_commit_appends_items_and_recalculates(
    db_session: Session, tenant: Tenant, rab: Rab
) -> None:
    job = ok(
        upload(
            rab,
            csv(
                "K-1;Semen 50 kg;Material;material;10;sak;65000;Toko A",
                ";Bongkar;;jasa;1;ls;;",
                ";Besi;;;12;batang;4.500;",
            ),
            name="../../harga lama.csv",
        ),
        201,
    )
    assert job["filename"] == "harga lama.csv"
    assert (job["row_count"], job["valid_count"], job["error_count"]) == (3, 2, 1)
    assert rab.version()["items"][-1]["name"] == "Item lama"  # nothing saved by preview

    committed = ok(rab.browser.post(rab.url(1, f"/imports/{job['id']}/commit")))
    assert committed["status"] == "committed"
    assert committed["result"] == {"items_created": 2}
    version = rab.version()
    names = [(i["position"], i["name"]) for i in version["items"]]
    assert names == [(1, "Item lama"), (2, "Semen 50 kg"), (3, "Bongkar")]
    semen, bongkar = version["items"][1], version["items"][2]
    assert semen["price_source_type"] == "import" and semen["price_source"] == "Toko A"
    assert semen["subtotal"] == "650000.00"
    assert bongkar["unit_price"] is None and bongkar["price_source_type"] is None
    assert Decimal(version["totals"]["subtotal"]) == Decimal("651000.00")
    assert version["totals"]["unpriced_items"] == 1

    # Retried commit: no duplicates. Discarding a committed job is refused.
    ok(rab.browser.post(rab.url(1, f"/imports/{job['id']}/commit")))
    assert len(rab.version()["items"]) == 3
    assert rab.browser.post(rab.url(1, f"/imports/{job['id']}/discard")).status_code == 409
    summary = db_session.execute(
        select(AuditLog.summary).where(AuditLog.action == "rab_item.import")
    ).scalar_one()
    assert (summary or "").endswith("2 item dari harga lama.csv")


def test_default_source_names_the_file_row(rab: Rab) -> None:
    job = ok(upload(rab, csv(";Pasir;;;2;m3;285000;")), 201)
    ok(rab.browser.post(rab.url(1, f"/imports/{job['id']}/commit")))
    assert rab.version()["items"][-1]["price_source"] == "Impor item.csv baris 2"


def test_only_draft_versions_accept_imports(rab: Rab) -> None:
    job = ok(upload(rab, csv(";Pasir;;;2;m3;285000;")), 201)
    ok(rab.act("submit"))
    assert upload(rab, csv(";Pasir;;;2;m3;285000;")).status_code == 409
    refused = rab.browser.post(rab.url(1, f"/imports/{job['id']}/commit"))
    assert refused.status_code == 409
    discarded = ok(rab.browser.post(rab.url(1, f"/imports/{job['id']}/discard")))
    assert discarded["status"] == "discarded"


def test_bad_files_and_permissions(
    db_session: Session, tenant: Tenant, other_tenant: Tenant, rab: Rab, login_as: LoginAs
) -> None:
    assert upload(rab, b"", name="kosong.csv").status_code == 422
    assert upload(rab, b"PK\x03\x04junk", name="item.xlsx").status_code == 422
    assert upload(rab, b"%PDF-1.7", name="item.pdf").status_code == 422
    assert upload(rab, b"x" * (5 * 1024 * 1024 + 1)).status_code == 413

    job = ok(upload(rab, csv(";Pasir;;;2;m3;285000;")), 201)
    viewer = login_as(make_user(db_session, tenant.org, Role.VIEWER))
    assert upload(rab, csv(";A;;;1;m;1;"), browser=viewer).status_code == 403
    assert viewer.post(rab.url(1, f"/imports/{job['id']}/commit")).status_code == 403
    assert viewer.get(rab.url(1, f"/imports/{job['id']}")).status_code == 200
    estimator = login_as(make_user(db_session, tenant.org, Role.ESTIMATOR))
    assert upload(rab, csv(";A;;;1;m;1;"), browser=estimator).status_code == 201

    # Jobs are bound to their version, kind and organization.
    assert rab.browser.get(f"/api/v1/price-catalog/imports/{job['id']}").status_code == 404
    assert other_tenant.browser.get(rab.url(1, f"/imports/{job['id']}")).status_code == 404
    assert other_tenant.browser.post(rab.url(1, f"/imports/{job['id']}/commit")).status_code == 404
