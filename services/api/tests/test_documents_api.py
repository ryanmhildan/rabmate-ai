"""Document API: preview/export, idempotent issuance, frozen snapshots, voiding, settings."""

from __future__ import annotations

import hashlib
import io
from decimal import Decimal
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import func, select, update
from sqlalchemy.orm import Session

from app.modules.audit.models import AuditLog
from app.modules.documents.models import IssuedDocument
from app.modules.documents.snapshot import snapshot_sha256
from app.modules.identity.permissions import Role
from app.modules.rab.models import RabVersion
from tests.factories import LoginAs, Tenant, make_client, make_project, make_user
from tests.test_rab_api import Rab, ok

KEY = "issue-key-0001"


def pdf_text(content: bytes) -> str:
    return " ".join(
        " ".join(page.extract_text().split()) for page in PdfReader(io.BytesIO(content)).pages
    )


@pytest.fixture
def rab(db_session: Session, tenant: Tenant) -> Rab:
    client = make_client(db_session, tenant.org, name="CV Klien Makmur", city="Jakarta")
    project = make_project(
        db_session,
        tenant.org,
        client_id=client.id,
        name="Instalasi Jaringan",
        tax_label="PPN",
        tax_rate_percent=Decimal("11"),
        currency="IDR",
    )
    created = ok(
        tenant.browser.post(f"/api/v1/projects/{project.id}/rab", json={"title": "RAB Jaringan"}),
        201,
    )
    rab = Rab(tenant.browser, created)
    ok(
        rab.add(
            name="Kabel UTP Cat6",
            category="Material",
            quantity="305",
            unit="m",
            unit_price="4300",
            specification="Belden",
        ),
        201,
    )
    ok(
        rab.add(
            name="Jasa instalasi",
            item_type="service",
            category="Jasa",
            quantity="48",
            unit="titik",
            unit_price="175000",
        ),
        201,
    )
    return rab


def approve(rab: Rab, version: int = 1) -> Any:
    ok(rab.act("submit", version))
    return ok(rab.act("approve", version, note="OK"))


def issue(browser: TestClient, rab: Rab, version: int = 1, key: str | None = KEY) -> Any:
    headers = {"Idempotency-Key": key} if key else {}
    return browser.post(rab.url(version, "/documents"), headers=headers)


# --------------------------------------------------------------------------- preview & export


def test_preview_and_export_show_live_draft(
    db_session: Session, tenant: Tenant, rab: Rab, login_as: LoginAs, other_tenant: Tenant
) -> None:
    response = tenant.browser.get(rab.url(1, "/preview.pdf"))
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].startswith("inline;")
    assert response.headers["cache-control"] == "no-store"
    text = pdf_text(response.content)
    assert "PRATINJAU" in text
    assert "Kabel UTP Cat6" in text
    assert "Status versi: Draft" in text
    # 9.711.500 + PPN 11% (no discount) = 10.779.765: the server's figure, printed verbatim.
    assert rab.version()["totals"]["grand_total"] == "10779765.00"
    assert "Total Rp 10.779.765" in text

    download = tenant.browser.get(rab.url(1, "/preview.pdf?download=true"))
    assert download.headers["content-disposition"].startswith("attachment;")

    export = tenant.browser.get(rab.url(1, "/export.xlsx"))
    assert export.status_code == 200
    assert export.headers["content-disposition"].endswith('_v1_draft.xlsx"')
    sheet = load_workbook(io.BytesIO(export.content))["RAB"]
    assert "Kabel UTP Cat6" in [c.value for row in sheet.iter_rows() for c in row]

    viewer = login_as(make_user(db_session, tenant.org, Role.VIEWER))
    assert viewer.get(rab.url(1, "/preview.pdf")).status_code == 200
    assert other_tenant.browser.get(rab.url(1, "/preview.pdf")).status_code == 404
    assert other_tenant.browser.get(rab.url(1, "/export.xlsx")).status_code == 404


def test_only_approved_versions_can_be_issued_by_permitted_roles(
    db_session: Session, tenant: Tenant, rab: Rab, login_as: LoginAs
) -> None:
    refused = issue(tenant.browser, rab)
    assert refused.status_code == 409
    assert refused.json()["error"]["details"] == {"status": "draft"}
    approve(rab)
    estimator = login_as(make_user(db_session, tenant.org, Role.ESTIMATOR))
    assert issue(estimator, rab).status_code == 403
    finance = login_as(make_user(db_session, tenant.org, Role.FINANCE))
    assert issue(finance, rab).status_code == 201


def test_issue_freezes_snapshot_and_stores_files(
    db_session: Session, tenant: Tenant, rab: Rab
) -> None:
    version = approve(rab)
    document = ok(issue(tenant.browser, rab), 201)
    assert document["number"] == f"{rab.number}/V1"
    assert document["status"] == "issued"
    assert Decimal(document["total_amount"]) == Decimal(version["grand_total"])
    assert document["issued_by"]["id"] == str(tenant.owner.id)
    assert [f["format"] for f in document["files"]] == ["pdf", "xlsx"]

    detail = ok(tenant.browser.get(f"/api/v1/documents/{document['id']}"))
    snapshot = detail["snapshot"]
    assert snapshot_sha256(snapshot) == detail["snapshot_sha256"]
    assert snapshot["document"]["number"] == document["number"]
    assert snapshot["client"]["name"] == "CV Klien Makmur"
    assert [g["title"] for g in snapshot["groups"]] == ["Material", "Jasa"]
    assert [g["subtotal"] for g in snapshot["groups"]] == ["1311500.00", "8400000.00"]
    assert snapshot["totals"]["grand_total"] == version["grand_total"]
    assert snapshot["totals"]["amount_in_words"] == (
        "Sepuluh juta tujuh ratus tujuh puluh sembilan ribu tujuh ratus enam puluh lima rupiah"
    )
    assert snapshot["rab"]["status_label"] == "Disetujui"
    assert snapshot["rab"]["approved_by"] == tenant.owner.full_name

    pdf_meta = document["files"][0]
    pdf = tenant.browser.get(f"/api/v1/documents/{document['id']}/files/pdf")
    assert hashlib.sha256(pdf.content).hexdigest() == pdf_meta["sha256"]
    assert pdf.headers["etag"] == f'"{pdf_meta["sha256"]}"'
    assert pdf.headers["content-disposition"] == f'inline; filename="{rab.number}_V1.pdf"'
    text = pdf_text(pdf.content)
    assert document["number"] in text
    assert f"Kode verifikasi: {detail['snapshot_sha256'][:16]}" in text
    assert "PRATINJAU" not in text
    xlsx = tenant.browser.get(f"/api/v1/documents/{document['id']}/files/xlsx")
    assert xlsx.headers["content-disposition"].startswith("attachment;")
    assert hashlib.sha256(xlsx.content).hexdigest() == document["files"][1]["sha256"]

    logged = db_session.execute(
        select(AuditLog).where(AuditLog.action == "document.issue")
    ).scalar_one()
    assert logged.entity_id == document["id"] or str(logged.entity_id) == document["id"]


def test_issuing_is_idempotent_and_single(tenant: Tenant, db_session: Session, rab: Rab) -> None:
    approve(rab)
    first = ok(issue(tenant.browser, rab), 201)
    retry = issue(tenant.browser, rab)
    assert retry.status_code == 200
    assert retry.headers["idempotent-replayed"] == "true"
    assert retry.json()["id"] == first["id"]

    other_key = issue(tenant.browser, rab, key="another-key-0002")
    assert other_key.status_code == 409
    assert other_key.json()["error"]["code"] == "already_issued"
    assert other_key.json()["error"]["details"]["document_id"] == first["id"]
    assert issue(tenant.browser, rab, key=None).status_code == 409
    count = db_session.execute(select(func.count()).select_from(IssuedDocument)).scalar_one()
    assert count == 1

    # The same key cannot be reused for another version.
    ok(rab.browser.post(f"/api/v1/rab/{rab.id}/versions", json={}), 201)
    approve(rab, 2)
    reused = issue(tenant.browser, rab, version=2)
    assert reused.status_code == 422
    assert reused.json()["error"]["code"] == "idempotency_key_reused"
    assert issue(tenant.browser, rab, version=2, key="bad key!").status_code == 422


def test_issued_document_never_changes_after_later_edits(
    tenant: Tenant, db_session: Session, rab: Rab
) -> None:
    approve(rab)
    document = ok(issue(tenant.browser, rab), 201)
    before_pdf = tenant.browser.get(f"/api/v1/documents/{document['id']}/files/pdf").content
    before = ok(tenant.browser.get(f"/api/v1/documents/{document['id']}"))

    ok(tenant.browser.patch("/api/v1/organizations/current", json={"legal_name": "PT Nama Baru"}))
    ok(tenant.browser.put("/api/v1/document-settings", json={"version": 0, "tagline": "Baru!"}))
    ok(rab.browser.post(f"/api/v1/rab/{rab.id}/versions", json={}), 201)
    v2 = rab.version(2)
    item = v2["items"][0]
    ok(
        rab.browser.put(
            rab.url(2, f"/items/{item['id']}"),
            json={
                **{k: item[k] for k in ("name", "quantity", "unit", "category")},
                "unit_price": "5000",
                "row_version": item["row_version"],
            },
        )
    )
    approve(rab, 2)

    after = ok(tenant.browser.get(f"/api/v1/documents/{document['id']}"))
    assert after["snapshot"] == before["snapshot"]
    assert after["snapshot_sha256"] == before["snapshot_sha256"]
    assert after["status"] == "issued"
    assert tenant.browser.get(f"/api/v1/documents/{document['id']}/files/pdf").content == before_pdf
    assert "PT Nama Baru" not in pdf_text(before_pdf)

    second = ok(issue(tenant.browser, rab, version=2, key="issue-key-v2"), 201)
    assert second["number"] == f"{rab.number}/V2"
    snapshot = ok(tenant.browser.get(f"/api/v1/documents/{second['id']}"))["snapshot"]
    assert snapshot["company"]["legal_name"] == "PT Nama Baru"
    assert snapshot["settings"]["tagline"] == "Baru!"
    assert snapshot["items"][0]["unit_price"] == "5000.0000"


def test_void_requires_reason_and_allows_reissue(
    tenant: Tenant, db_session: Session, rab: Rab, login_as: LoginAs
) -> None:
    approve(rab)
    document = ok(issue(tenant.browser, rab), 201)
    url = f"/api/v1/documents/{document['id']}/void"
    viewer = login_as(make_user(db_session, tenant.org, Role.VIEWER))
    assert viewer.post(url, json={"reason": "x"}).status_code == 403
    assert tenant.browser.post(url, json={"reason": "  "}).status_code == 422
    voided = ok(tenant.browser.post(url, json={"reason": "Salah klien"}))
    assert voided["status"] == "void"
    assert voided["void_reason"] == "Salah klien"
    assert voided["voided_by"]["id"] == str(tenant.owner.id)
    assert tenant.browser.post(url, json={"reason": "lagi"}).status_code == 409
    # Files of a voided document stay downloadable as issued.
    assert tenant.browser.get(f"/api/v1/documents/{document['id']}/files/pdf").status_code == 200

    summary = ok(tenant.browser.get("/api/v1/dashboard/summary"))["documents"]
    assert summary == {"issued": 0, "void": 1, "approved_rab_not_issued": 1}
    reissued = ok(issue(tenant.browser, rab, key="issue-key-again"), 201)
    summary = ok(tenant.browser.get("/api/v1/dashboard/summary"))["documents"]
    assert summary == {"issued": 1, "void": 1, "approved_rab_not_issued": 0}
    assert reissued["number"] == f"{rab.number}/V1-R1"
    actions = (
        db_session.execute(
            select(AuditLog.action)
            .where(AuditLog.entity_type == "document")
            .order_by(AuditLog.created_at)
        )
        .scalars()
        .all()
    )
    assert actions == ["document.issue", "document.void", "document.issue"]


def test_document_settings(
    tenant: Tenant, db_session: Session, rab: Rab, login_as: LoginAs
) -> None:
    defaults = ok(tenant.browser.get("/api/v1/document-settings"))
    assert defaults["version"] == 0
    assert defaults["accent_color"] == "#1F4E79"
    assert defaults["show_item_codes"] is True

    body = {
        "version": 0,
        "tagline": "Solusi jaringan",
        "accent_color": "#0F766E",
        "rab_terms": "Harga berlaku 30 hari.",
        "show_item_codes": False,
        "payment_bank_name": "Bank Contoh",
        "payment_account_number": "123-456-7890",
        "payment_account_holder": "PT Uji",
    }
    saved = ok(tenant.browser.put("/api/v1/document-settings", json=body))
    assert saved["version"] == 1
    assert saved["payment_account_number"] == "123-456-7890"
    assert tenant.browser.put("/api/v1/document-settings", json=body).status_code == 409
    assert (
        tenant.browser.put(
            "/api/v1/document-settings", json={**body, "version": 1, "accent_color": "blue"}
        ).status_code
        == 422
    )
    assert (
        tenant.browser.put(
            "/api/v1/document-settings",
            json={**body, "version": 1, "payment_account_number": "12ab"},
        ).status_code
        == 422
    )
    manager = login_as(make_user(db_session, tenant.org, Role.PROJECT_MANAGER))
    assert manager.get("/api/v1/document-settings").status_code == 200
    assert manager.put("/api/v1/document-settings", json={**body, "version": 1}).status_code == 403

    text = pdf_text(tenant.browser.get(rab.url(1, "/preview.pdf")).content)
    assert "Solusi jaringan" in text
    assert "Harga berlaku 30 hari." in text
    assert "Kode" not in text.split("Uraian")[0][-20:]  # no code column before "Uraian"
    assert (
        db_session.execute(
            select(func.count())
            .select_from(AuditLog)
            .where(AuditLog.action == "document_settings.update")
        ).scalar_one()
        == 1
    )


def test_list_filters_and_isolation(tenant: Tenant, other_tenant: Tenant, rab: Rab) -> None:
    approve(rab)
    document = ok(issue(tenant.browser, rab), 201)
    listed = ok(tenant.browser.get(f"/api/v1/documents?rab_id={rab.id}"))
    assert [d["id"] for d in listed["items"]] == [document["id"]]
    assert ok(tenant.browser.get("/api/v1/documents?status=void"))["total"] == 0
    assert ok(tenant.browser.get(f"/api/v1/documents?q={rab.number}"))["total"] == 1
    assert ok(tenant.browser.get("/api/v1/documents?q=tidak-ada"))["total"] == 0

    stranger = other_tenant.browser
    assert ok(stranger.get("/api/v1/documents"))["total"] == 0
    assert stranger.get(f"/api/v1/documents/{document['id']}").status_code == 404
    assert stranger.get(f"/api/v1/documents/{document['id']}/files/pdf").status_code == 404
    assert (
        stranger.post(f"/api/v1/documents/{document['id']}/void", json={"reason": "x"}).status_code
        == 404
    )
    assert issue(stranger, rab, key="stranger-key-01").status_code == 404


def test_inconsistent_stored_totals_block_documents(
    tenant: Tenant, db_session: Session, rab: Rab
) -> None:
    db_session.execute(
        update(RabVersion).where(RabVersion.rab_id == rab.id).values(grand_total=Decimal("1.00"))
    )
    db_session.commit()
    response = tenant.browser.get(rab.url(1, "/preview.pdf"))
    assert response.status_code == 409
    assert "tidak sesuai" in response.json()["error"]["message"]
