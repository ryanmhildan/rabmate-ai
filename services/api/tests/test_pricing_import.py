"""Spreadsheet import: file safety, row validation, preview/commit, template."""

from __future__ import annotations

import io
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.modules.pricing.importer import ImportFileError, parse_file
from app.modules.pricing.imports import safe_filename
from app.modules.pricing.models import PriceCatalogItem, PriceObservation, Supplier
from tests.factories import Tenant

HEADER = "kode;nama;satuan;harga;tanggal_harga;supplier;merek"


def csv_bytes(*lines: str) -> bytes:
    return ("\n".join(lines) + "\n").encode("utf-8")


def xlsx_bytes(rows: list[list[Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def upload(browser: TestClient, content: bytes, filename: str = "harga.csv") -> Any:
    return browser.post(
        "/api/v1/price-catalog/imports",
        files={"file": (filename, content, "application/octet-stream")},
    )


# --------------------------------------------------------------------------- parsing (no DB)


def test_csv_with_semicolons_and_indonesian_numbers() -> None:
    parsed = parse_file(
        "harga.csv",
        csv_bytes(
            HEADER,
            "NET-01;Kabel UTP Cat6;m;4.500,75;19/10/2026;Toko Jaya;Belden",
            ";Semen 50kg;sak;65000;2026-10-01;;",
        ),
    )
    first, second = parsed.rows
    assert (
        first.valid
        and first.data["price"] == "4500.75"
        and first.data["observed_at"] == "2026-10-19"
    )
    assert second.valid and second.data["code"] is None and second.data["item_type"] == "material"


def test_comma_csv_and_header_aliases_with_unknown_column_warning() -> None:
    parsed = parse_file(
        "harga.csv",
        csv_bytes("Nama Item,Unit,Harga Satuan,Tanggal,Warna", "Pasir,m3,350000,2026-10-01,merah"),
    )
    assert parsed.rows[0].valid
    assert any("Warna" in w for w in parsed.warnings)


def test_row_errors_are_specific_and_never_guessed() -> None:
    parsed = parse_file(
        "harga.csv",
        csv_bytes(
            HEADER,
            "A;Kabel;m;4.500;2026-10-01;;",  # ambiguous thousands
            "B;;m;100;2026-10-01;;",  # missing name
            "C;Pipa;m;abc;2026-13-01;;",  # bad price and date
            "D;Baut;pcs;-5;2026-10-01;;",  # negative
            "E;Kabel;m;100;2026-10-01;;",
            "E;Kabel;m;100;2026-10-01;;",  # duplicate of previous row
        ),
    )
    errors = {row.row: row.errors for row in parsed.rows}
    assert any("ambigu" in e for e in errors[2])
    assert any("Nama" in e for e in errors[3])
    assert any("Harga" in e for e in errors[4]) and any("tanggal" in e.lower() for e in errors[4])
    assert any("negatif" in e for e in errors[5])
    assert errors[6] == []
    assert any("Duplikat dari baris 6" in e for e in errors[7])


@pytest.mark.parametrize(
    ("filename", "content", "message"),
    [
        ("harga.txt", b"nama;satuan", "Format"),
        ("harga.csv", b"", "kosong"),
        ("harga.xlsx", b"not a zip", "bukan XLSX"),
        ("harga.csv", csv_bytes("nama;satuan", "x;m"), "Kolom wajib"),
        ("harga.csv", csv_bytes("nama;nama;satuan;harga;tanggal_harga"), "lebih dari sekali"),
        ("harga.csv", csv_bytes("nama;satuan;harga;tanggal_harga"), "tidak berisi baris"),
    ],
)
def test_file_level_rejections(filename: str, content: bytes, message: str) -> None:
    with pytest.raises(ImportFileError, match=message):
        parse_file(filename, content)


def test_xlsx_formula_cells_are_rejected_not_evaluated() -> None:
    content = xlsx_bytes(
        [
            ["nama", "satuan", "harga", "tanggal_harga"],
            ["Kabel", "m", 4500.75, "2026-10-01"],
            ["Semen", "sak", "=A1*1000", "2026-10-01"],
        ]
    )
    rows = parse_file("harga.xlsx", content).rows
    assert rows[0].valid and rows[0].data["price"] == "4500.75"
    assert not rows[1].valid
    assert any("rumus" in e for e in rows[1].errors)


def test_xlsx_dates_and_numeric_codes() -> None:
    from datetime import datetime

    content = xlsx_bytes(
        [
            ["kode", "nama", "satuan", "harga", "tanggal_harga"],
            [1001, "Kabel", "m", 4500, datetime(2026, 10, 19)],
        ]
    )
    row = parse_file("harga.xlsx", content).rows[0]
    assert row.valid and row.data["code"] == "1001" and row.data["observed_at"] == "2026-10-19"


# --------------------------------------------------------------------------- API flow


def test_preview_then_commit_creates_unverified_prices(db_session: Session, tenant: Tenant) -> None:
    existing = tenant.browser.post(
        "/api/v1/price-catalog", json={"code": "NET-01", "name": "Kabel UTP Cat6", "unit": "m"}
    ).json()
    content = csv_bytes(
        HEADER,
        "NET-01;Kabel UTP Cat6;m;4500;2026-10-19;Toko Jaya;",  # adds a price to the existing item
        ";Semen 50kg;sak;65000;2026-10-01;Toko Jaya;Tiga Roda",  # new item, same new supplier
        ";Semen 50kg;sak;66000;2026-10-15;;Tiga Roda",  # same new item again → one item
        "NET-01;Kabel;roll;1;2026-10-01;;",  # code exists with another unit → error
        ";Pipa;m;abc;2026-10-01;;",  # invalid
    )

    preview = upload(tenant.browser, content)
    assert preview.status_code == 201, preview.text
    job = preview.json()
    assert (job["status"], job["row_count"], job["valid_count"], job["error_count"]) == (
        "previewed",
        5,
        3,
        2,
    )
    actions = [r["action"] for r in job["rows"]]
    assert actions[:3] == ["add_price", "create_item", "create_item"]
    # Nothing is written until commit.
    assert db_session.execute(select(func.count()).select_from(PriceObservation)).scalar_one() == 0

    committed = tenant.browser.post(f"/api/v1/price-catalog/imports/{job['id']}/commit").json()
    assert committed["status"] == "committed"
    assert committed["result"] == {"items_created": 1, "prices_created": 3, "suppliers_created": 1}

    observations = db_session.execute(select(PriceObservation)).unique().scalars().all()
    assert {o.verification_status.value for o in observations} == {"unverified"}
    assert {o.source_type.value for o in observations} == {"import"}
    assert all(o.source_reference for o in observations)  # provenance points at file + row
    semen = db_session.execute(
        select(PriceCatalogItem).where(PriceCatalogItem.name == "Semen 50kg")
    ).scalar_one()
    assert semen.brand == "Tiga Roda"
    assert db_session.execute(select(func.count()).select_from(Supplier)).scalar_one() == 1
    detail = tenant.browser.get(f"/api/v1/price-catalog/{existing['id']}").json()
    assert detail["prices"]["observation_count"] == 1


def test_commit_is_idempotent_and_discard_is_final(db_session: Session, tenant: Tenant) -> None:
    job = upload(tenant.browser, csv_bytes(HEADER, ";Bata merah;bh;900;2026-10-01;;")).json()
    first = tenant.browser.post(f"/api/v1/price-catalog/imports/{job['id']}/commit")
    second = tenant.browser.post(f"/api/v1/price-catalog/imports/{job['id']}/commit")
    assert first.status_code == second.status_code == 200
    assert db_session.execute(select(func.count()).select_from(PriceObservation)).scalar_one() == 1
    assert (
        tenant.browser.post(f"/api/v1/price-catalog/imports/{job['id']}/discard").status_code == 409
    )

    other = upload(tenant.browser, csv_bytes(HEADER, ";Genteng;bh;5000;2026-10-01;;")).json()
    assert (
        tenant.browser.post(f"/api/v1/price-catalog/imports/{other['id']}/discard").json()["status"]
        == "discarded"
    )
    assert (
        tenant.browser.post(f"/api/v1/price-catalog/imports/{other['id']}/commit").status_code
        == 409
    )


def test_upload_rejections_and_isolation(tenant: Tenant, other_tenant: Tenant) -> None:
    assert upload(tenant.browser, b"x" * (5 * 1024 * 1024 + 1)).status_code == 413
    bad = upload(tenant.browser, b"nama;satuan\nx;m\n")
    assert bad.status_code == 422 and "Kolom wajib" in bad.json()["error"]["message"]
    only_errors = upload(tenant.browser, csv_bytes(HEADER, ";;m;1;2026-10-01;;")).json()
    assert (
        tenant.browser.post(f"/api/v1/price-catalog/imports/{only_errors['id']}/commit").status_code
        == 422
    )

    job = upload(tenant.browser, csv_bytes(HEADER, ";Kabel;m;1;2026-10-01;;")).json()
    assert other_tenant.browser.get(f"/api/v1/price-catalog/imports/{job['id']}").status_code == 404
    assert (
        other_tenant.browser.post(f"/api/v1/price-catalog/imports/{job['id']}/commit").status_code
        == 404
    )


def test_filename_is_sanitized_for_display(tenant: Tenant) -> None:
    job = upload(
        tenant.browser, csv_bytes(HEADER, ";Kabel;m;1;2026-10-01;;"), "../../evil.csv"
    ).json()
    assert job["filename"] == "evil.csv"
    # Control characters and Windows paths are stripped too (HTTP clients usually encode them).
    assert safe_filename("C:\\tmp\\a\x07b.csv") == "ab.csv"
    assert safe_filename(None) == "upload"


def test_templates_download_and_reimport(tenant: Tenant) -> None:
    csv_template = tenant.browser.get(
        "/api/v1/price-catalog/imports/template", params={"format": "csv"}
    )
    xlsx_template = tenant.browser.get("/api/v1/price-catalog/imports/template")
    assert csv_template.status_code == xlsx_template.status_code == 200
    assert "attachment" in xlsx_template.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(xlsx_template.content))
    assert workbook.sheetnames == ["Katalog", "Petunjuk"]
    headers = [c.value for c in next(workbook["Katalog"].iter_rows())]
    assert headers[:2] == ["kode", "nama"]
    # The empty template is a valid header row (but has no data rows).
    with pytest.raises(ImportFileError, match="tidak berisi baris"):
        parse_file("t.xlsx", xlsx_template.content)
