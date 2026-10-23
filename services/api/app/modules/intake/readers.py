"""Turn an uploaded file into pages of text lines for extraction (ADR-043).

- PDF: the text layer, in pypdf's layout mode so table columns survive; runs of two or more
  spaces become " | ". A page without text but with an image is treated as a scan and goes to
  OCR when available, otherwise it is reported as unreadable.
- XLSX: every sheet, row by row, cells joined with " | ". Formulas are never evaluated (shown
  as "(rumus)"); numbers are written in Indonesian notation (1.234,5).
- CSV: rows as written (the user picks the number format).
- PNG/JPEG: OCR when available.

Limits keep a hostile or huge file from exhausting the server: pages, sheets, rows, characters,
image pixels, OCR pages. Everything read here is untrusted text for the extraction rules.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from PIL import Image, UnidentifiedImageError

from app.modules.documents import formatting as fmt
from app.modules.intake.models import SourceKind
from app.modules.intake.ocr import OcrEngine, OcrError
from app.modules.intake.sniff import UnsupportedFileError

logger = logging.getLogger(__name__)

MAX_LINE_CHARS = 500
MAX_TOTAL_CHARS = 300_000
MAX_SHEETS = 10
MAX_ROWS_PER_SHEET = 3000
MAX_CELLS_PER_ROW = 40
MAX_IMAGE_PIXELS = 40_000_000
MIN_PAGE_TEXT = 20
_GAP = re.compile(r" {2,}")


@dataclass
class Page:
    no: int
    label: str
    method: str  # "teks", "tabel", "ocr", "kosong"
    lines: list[str] = field(default_factory=list)


@dataclass
class Extraction:
    pages: list[Page]
    warnings: list[str]
    ocr: dict[str, Any]

    def to_json(self) -> dict[str, Any]:
        return {
            "pages": [
                {"no": p.no, "label": p.label, "method": p.method, "lines": p.lines}
                for p in self.pages
            ],
            "warnings": self.warnings,
            "ocr": self.ocr,
        }


def normalize_lines(text: str) -> list[str]:
    lines = []
    for raw in text.replace("\t", "   ").splitlines():
        line = _GAP.sub(" | ", raw.strip())
        if line:
            lines.append(line[:MAX_LINE_CHARS])
    return lines


class _Budget:
    def __init__(self, warnings: list[str]) -> None:
        self.chars = 0
        self.warnings = warnings
        self.exhausted = False

    def take(self, lines: list[str]) -> list[str]:
        kept = []
        for line in lines:
            if self.chars + len(line) > MAX_TOTAL_CHARS:
                if not self.exhausted:
                    limit = fmt.number(Decimal(MAX_TOTAL_CHARS), max_places=0)
                    self.warnings.append(
                        f"Dokumen melebihi {limit} karakter; sisanya tidak dibaca."
                    )
                self.exhausted = True
                break
            self.chars += len(line)
            kept.append(line)
        return kept


def _png(image: Image.Image) -> bytes:
    if image.width * image.height > MAX_IMAGE_PIXELS:
        raise OcrError("Gambar terlalu besar untuk OCR.")
    if image.mode not in ("1", "L", "LA", "P", "RGB", "RGBA"):
        image = image.convert("RGB")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


# --------------------------------------------------------------------------- PDF


def _read_pdf(content: bytes, ocr: OcrEngine, max_pages: int, ocr_max_pages: int) -> Extraction:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    warnings: list[str] = []
    locked = UnsupportedFileError("PDF dikunci kata sandi; buka kuncinya lalu unggah ulang.")
    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted and not reader.decrypt(""):
            raise locked
        total = len(reader.pages)
    except UnsupportedFileError:
        raise
    except (PyPdfError, ValueError, KeyError, OSError) as exc:
        if "encrypt" in str(exc).lower() or "decrypt" in str(exc).lower():
            raise locked from None
        raise UnsupportedFileError(f"PDF tidak dapat dibaca ({type(exc).__name__}).") from None
    if total > max_pages:
        warnings.append(f"Hanya {max_pages} halaman pertama dari {total} yang dibaca.")
    budget = _Budget(warnings)
    status = ocr.status()
    ocr_used = 0
    pages: list[Page] = []
    for index in range(min(total, max_pages)):
        no = index + 1
        try:
            page = reader.pages[index]
            text = page.extract_text(extraction_mode="layout") or ""
        except Exception as exc:  # a broken page must not sink the whole document
            logger.info("pdf page %d unreadable: %s", no, type(exc).__name__)
            warnings.append(f"Halaman {no} tidak dapat dibaca.")
            continue
        lines = normalize_lines(text)
        method = "teks"
        if sum(len(line) for line in lines) < MIN_PAGE_TEXT:
            lines, method = _ocr_page(
                page, no, ocr, status.available, ocr_used, ocr_max_pages, warnings
            )
            if method == "ocr":
                ocr_used += 1
        pages.append(Page(no, f"Halaman {no}", method, budget.take(lines)))
        if budget.exhausted:
            break
    return Extraction(pages, warnings, {"available": status.available, "message": status.message})


def _ocr_page(
    page: Any,
    no: int,
    ocr: OcrEngine,
    available: bool,
    used: int,
    limit: int,
    warnings: list[str],
) -> tuple[list[str], str]:
    try:
        images = list(page.images)
    except Exception:
        images = []
    if not images:
        return [], "kosong"
    if not available:
        warnings.append(f"Halaman {no} tampaknya hasil scan; OCR tidak tersedia, dilewati.")
        return [], "kosong"
    if used >= limit:
        warnings.append(f"Halaman {no} hasil scan tidak di-OCR (batas {limit} halaman).")
        return [], "kosong"
    try:
        largest = max(images, key=lambda img: len(img.data))
        text = ocr.image_to_text(_png(largest.image))
    except (OcrError, OSError, ValueError) as exc:
        warnings.append(f"Halaman {no}: OCR gagal ({exc}).")
        return [], "kosong"
    return normalize_lines(text), "ocr"


# --------------------------------------------------------------------------- spreadsheets


def _cell_text(cell: Any) -> str:
    value = cell.value
    if value is None:
        return ""
    if getattr(cell, "data_type", None) == "f" or (
        isinstance(value, str) and value.startswith("=")
    ):
        return "(rumus)"
    if isinstance(value, bool):
        return "ya" if value else "tidak"
    if isinstance(value, int):
        return fmt.number(Decimal(value), max_places=0)
    if isinstance(value, float):
        return fmt.number(Decimal(repr(value)), max_places=4)
    if isinstance(value, Decimal):
        return fmt.number(value, max_places=4)
    if isinstance(value, datetime | date):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return " ".join(str(value).split())


def _row_line(cells: list[str]) -> str:
    while cells and not cells[-1]:
        cells.pop()
    return " | ".join(cells)[:MAX_LINE_CHARS] if any(cells) else ""


def _read_xlsx(content: bytes) -> Extraction:
    import zipfile

    from openpyxl import load_workbook

    warnings: list[str] = []
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise UnsupportedFileError(f"XLSX tidak dapat dibaca ({type(exc).__name__}).") from None
    budget = _Budget(warnings)
    pages: list[Page] = []
    try:
        sheets = workbook.worksheets
        if len(sheets) > MAX_SHEETS:
            warnings.append(f"Hanya {MAX_SHEETS} sheet pertama yang dibaca.")
        for no, sheet in enumerate(sheets[:MAX_SHEETS], start=1):
            lines: list[str] = []
            for count, row in enumerate(sheet.iter_rows()):
                if count >= MAX_ROWS_PER_SHEET:
                    warnings.append(
                        f"Sheet '{sheet.title}': hanya {MAX_ROWS_PER_SHEET} baris dibaca."
                    )
                    break
                line = _row_line([_cell_text(c) for c in row[:MAX_CELLS_PER_ROW]])
                if line:
                    lines.append(line)
            pages.append(Page(no, f"Sheet {sheet.title}", "tabel", budget.take(lines)))
            if budget.exhausted:
                break
    finally:
        workbook.close()
    return Extraction(pages, warnings, {})


def _read_csv(content: bytes) -> Extraction:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252", errors="replace")
    sample = text[:4096]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    warnings: list[str] = []
    budget = _Budget(warnings)
    lines: list[str] = []
    for count, record in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter)):
        if count >= MAX_ROWS_PER_SHEET:
            warnings.append(f"Hanya {MAX_ROWS_PER_SHEET} baris pertama yang dibaca.")
            break
        line = _row_line([" ".join(v.split()) for v in record[:MAX_CELLS_PER_ROW]])
        if line:
            lines.append(line)
    return Extraction([Page(1, "CSV", "tabel", budget.take(lines))], warnings, {})


# --------------------------------------------------------------------------- images


def _read_image(content: bytes, ocr: OcrEngine) -> Extraction:
    status = ocr.status()
    meta = {"available": status.available, "message": status.message}
    try:
        with Image.open(io.BytesIO(content)) as probe:
            probe.verify()
        with Image.open(io.BytesIO(content)) as image:
            image.load()
            png = _png(image)
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
        raise UnsupportedFileError(f"Gambar tidak dapat dibaca ({type(exc).__name__}).") from None
    except OcrError as exc:
        raise UnsupportedFileError(str(exc)) from None
    if not status.available:
        return Extraction([], [f"Gambar hanya dapat dibaca dengan OCR. {status.message}"], meta)
    try:
        lines = normalize_lines(ocr.image_to_text(png))
    except OcrError as exc:
        return Extraction([], [f"OCR gagal: {exc}"], meta)
    return Extraction([Page(1, "Gambar", "ocr", lines[:MAX_ROWS_PER_SHEET])], [], meta)


def read_document(
    kind: SourceKind,
    content: bytes,
    *,
    ocr: OcrEngine,
    max_pages: int,
    ocr_max_pages: int,
) -> Extraction:
    if kind is SourceKind.PDF:
        return _read_pdf(content, ocr, max_pages, ocr_max_pages)
    if kind is SourceKind.XLSX:
        return _read_xlsx(content)
    if kind is SourceKind.CSV:
        return _read_csv(content)
    return _read_image(content, ocr)
